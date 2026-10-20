#!/usr/bin/env python3
"""Synthetic DOCX/PDF smoke in the existing restricted analysis-container boundary.

No case/evidence input, model calls, network, package installs, or tag changes.
Run with an explicit candidate image and a new empty local output directory.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import uuid


def inside():
    import docx
    import lxml.etree
    from docx.oxml.ns import qn
    from docx.shared import Pt

    assert os.getuid() != 0
    root = Path('/work')
    title = '포식이 문서 생성 검증'
    sentence = '합성 자료로 Word 문서와 PDF 변환을 확인합니다.'
    english = 'Synthetic document creation and PDF conversion verified.'
    document = docx.Document()
    for style_name in ('Normal', 'Title'):
        style = document.styles[style_name]
        style.font.name = 'Noto Sans CJK KR'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK KR')
    document.styles['Normal'].font.size = Pt(11)
    document.add_paragraph(title, 'Title')
    document.add_paragraph(sentence)
    document.add_paragraph(english)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '검사'
    table.cell(0, 1).text = '결과'
    table.cell(1, 0).text = '한국어와 English'
    table.cell(1, 1).text = '정상'
    source = root / 'synthetic.docx'
    document.save(source)
    reopened = docx.Document(source)
    assert sentence in '\n'.join(p.text for p in reopened.paragraphs)
    conversion = subprocess.run([
        'libreoffice', '-env:UserInstallation=file:///tmp/forsic-document-smoke-profile',
        '--headless', '--convert-to', 'pdf', '--outdir', str(root), str(source),
    ], check=True, capture_output=True, text=True, timeout=90)
    pdf = root / 'synthetic.pdf'
    assert pdf.is_file() and pdf.read_bytes().startswith(b'%PDF-')
    subprocess.run(['pdftotext', '-layout', str(pdf), str(root / 'synthetic.txt')], check=True)
    extracted = (root / 'synthetic.txt').read_text()
    for expected in (title, sentence, english, '한국어와 English'):
        assert expected in extracted, 'PDF text round-trip mismatch: '+expected
    subprocess.run(['pdftoppm', '-singlefile', '-scale-to', '1400', '-png', str(pdf),
                    str(root / 'synthetic-page')], check=True)
    packages = subprocess.check_output(['dpkg-query', '-W', '-f=${Package}\t${Version}\n'], text=True)
    (root / 'packages.tsv').write_text(packages)
    record = {
        'status': 'passed', 'synthetic_only': True, 'uid': os.getuid(),
        'python': sys.version.split()[0], 'python_docx': docx.__version__,
        'lxml': lxml.etree.LXML_VERSION,
        'libreoffice': subprocess.check_output(['libreoffice', '--version'], text=True).strip(),
        'korean_and_english_pdf_text_verified': True,
        'conversion_stdout': conversion.stdout.strip(),
        'conversion_stderr': conversion.stderr.strip(),
        'files': {p.name: {'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
                  for p in (source, pdf, root / 'synthetic.txt', root / 'synthetic-page.png')},
    }
    (root / 'container-verification.json').write_text(json.dumps(record, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(record, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--image')
    parser.add_argument('--output', type=Path)
    parser.add_argument('--inside', action='store_true', help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.inside:
        inside()
        return
    if not args.image or not args.output:
        parser.error('--image and --output are required')
    if os.getuid() == 0:
        parser.error('Use the non-root analysis-shell operator')
    output = args.output.resolve()
    if output.exists() and any(output.iterdir()):
        parser.error('Output directory must be new or empty')
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    image = json.loads(subprocess.check_output(['docker', 'image', 'inspect', args.image], text=True))[0]
    container_name = 'forsic-doc-smoke-'+uuid.uuid4().hex
    command = ['docker', 'run', '--rm', '--pull=never', '--name', container_name,
               '--init', '--network', 'none', '--read-only', '--cap-drop', 'ALL',
               '--security-opt', 'no-new-privileges', '--user', f'{os.getuid()}:{os.getgid()}',
               '--memory', '8g', '--memory-swap', '8g', '--cpus', '2', '--pids-limit', '64',
               '--ulimit', 'fsize=268435456:268435456', '--tmpfs', '/tmp:rw,nosuid,nodev,size=256m',
               '--env', 'HOME=/work', '--env', 'TMPDIR=/tmp', '--workdir', '/work',
               '--mount', f'type=bind,src={output},dst=/work',
               '--mount', f'type=bind,src={Path(__file__).resolve()},dst=/opt/smoke_document_runtime.py,readonly,bind-propagation=rprivate,bind-recursive=readonly',
               image['Id'], 'python', '-B', '/opt/smoke_document_runtime.py', '--inside']
    receipt = {'image_tag': args.image, 'image_id': image['Id'], 'image_size_bytes': image['Size'],
               'command': command}
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=120)
        receipt.update(exit_code=result.returncode, stdout=result.stdout, stderr=result.stderr)
    except subprocess.TimeoutExpired as exc:
        receipt.update(exit_code=124, timed_out=True,
                       stdout=(exc.stdout or b'').decode(errors='replace') if isinstance(exc.stdout, bytes) else exc.stdout or '',
                       stderr=(exc.stderr or b'').decode(errors='replace') if isinstance(exc.stderr, bytes) else exc.stderr or '')
    except OSError as exc:
        receipt.update(exit_code=1, stdout='', stderr=str(exc))
    finally:
        # Killing a Docker client does not kill its container. The name is unique
        # to this synthetic run, so cleanup cannot target investigation work.
        try:
            cleanup = subprocess.run(['docker', 'rm', '-f', container_name],
                                     capture_output=True, text=True, timeout=15)
            receipt['cleanup'] = {'exit_code': cleanup.returncode,
                                  'stdout': cleanup.stdout, 'stderr': cleanup.stderr}
        except (OSError, subprocess.TimeoutExpired) as exc:
            receipt['cleanup'] = {'error': str(exc)}
    (output / 'runtime-smoke.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k: receipt[k] for k in ('image_tag', 'image_id', 'image_size_bytes', 'exit_code')}))
    if receipt['exit_code']:
        print(receipt['stderr'], file=sys.stderr)
        raise SystemExit(receipt['exit_code'])


if __name__ == '__main__':
    main()
