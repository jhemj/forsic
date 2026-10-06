"""Two reader views of one immutable investigation snapshot. No model calls."""
from datetime import datetime
from html import escape
import hashlib
import json
from pathlib import Path
import re
from zoneinfo import ZoneInfo

from .notes import current_notes, sources, timeline


def _text_list(value):
    return [str(x) for x in value] if isinstance(value, list) else [str(value)] if value else []


def prepare(case, args):
    if args.get('snapshot_id'):
        sid = args['snapshot_id']
        if not re.fullmatch(r'[a-f0-9]{64}', sid):
            raise ValueError('Invalid snapshot_id')
        path = case.output / ('snapshot-' + sid + '.json')
        raw = path.read_bytes()
        if hashlib.sha256(raw).hexdigest() != sid:
            raise ValueError('Snapshot content changed')
        return json.loads(raw), sid
    body = dict(args.get('snapshot') or {})
    if not body:
        notes = current_notes(case)
        body = {'title': case.config.get('label', '조사 보고서'),
                'question': case.config.get('question', ''),
                'summary': '\n\n'.join(n['answer'] for n in notes if n.get('evidence_ids')) or '아직 근거에 연결한 답변이 없습니다.',
                'findings': [{'title': n['question'], 'detail': n['answer'], 'evidence_ids': n['evidence_ids']} for n in notes if n.get('evidence_ids')],
                'alternatives': [a for n in notes for a in _text_list(n.get('alternatives'))],
                'critical_gaps': [a for n in notes for a in _text_list(n.get('critical_gaps'))],
                'gaps': [a for n in notes for a in _text_list(n.get('gaps'))] + [f"근거 미연결: {n['question']} — {n['answer']}" for n in notes if not n.get('evidence_ids')] or (['아직 검토한 질문이 없습니다.'] if not notes else []),
                'actions': [a for n in notes for a in _text_list(n.get('next_checks'))],
                'timeline': [a for n in notes for a in n.get('timeline', [])]}
    if not body.get('summary'):
        raise ValueError('snapshot.summary is required')
    if len(json.dumps(body, ensure_ascii=False)) > 100000:
        raise ValueError('Snapshot exceeds 100,000 characters; keep raw evidence in the source appendix')
    ids = []
    for finding in body.get('findings', []):
        if not finding.get('title') or not finding.get('detail') or not finding.get('evidence_ids'):
            raise ValueError('Each finding needs title, detail and actual evidence_ids; put unverified ideas in alternatives or gaps')
        ids.extend(finding['evidence_ids'])
    for item in body.get('timeline', []):
        if not item.get('evidence_ids'):
            raise ValueError('A timeline entry requires evidence_ids')
        ids.extend(item['evidence_ids'])
    records = sources(case, ids)
    # A prose summary must not silently discard the gaps/alternatives already
    # recorded against its evidence. Reuse case notes, not another model pass.
    cited_ids = set(ids)
    linked_notes = [n for n in current_notes(case) if cited_ids.intersection(n.get('evidence_ids', []))]
    for field, note_field in (('alternatives', 'alternatives'), ('gaps', 'gaps'), ('critical_gaps', 'critical_gaps'), ('actions', 'next_checks')):
        values = _text_list(body.get(field))
        values.extend(value for note in linked_notes for value in _text_list(note.get(note_field)))
        body[field] = list(dict.fromkeys(values))
    for record in records:
        request = case.event(record['data'].get('started_id', ''))
        record['request'] = request if request and request['kind'] == 'tool_start' else None
    body['timeline'] = timeline(body.get('timeline', []), case)
    body['case'] = case.info({})
    body['sources'] = records
    body['created_at'] = datetime.now(ZoneInfo('Asia/Seoul')).isoformat(timespec='seconds')
    body['analysis_status'] = 'partial' if body.get('critical_gaps') else args.get('analysis_status', 'partial')
    raw = json.dumps(body, ensure_ascii=False, sort_keys=True, indent=2).encode()
    sid = hashlib.sha256(raw).hexdigest()
    (case.output / ('snapshot-' + sid + '.json')).write_bytes(raw)
    return body, sid


def sections(body, audience):
    source_map = {r['id']: i + 1 for i, r in enumerate(body['sources'])}
    def cited(text, ids):
        return text + ' ' + ' '.join(f'[{source_map[eid]}]' for eid in dict.fromkeys(ids))
    result = [('조사 결론', _text_list(body['summary'])),
              ('조사 범위', [str(body.get('question') or body['case'].get('question', '')), str(body['case'].get('scope') or '')])]
    for finding in body.get('findings', []):
        label = {'observation': '확인 사실 · ', 'hypothesis': '가설 · '}.get(finding.get('kind'), '')
        result.append((label + finding['title'], [cited(finding['detail'], finding['evidence_ids'])]))
    for title, field in [('판단에 필요한 추가 확인', 'critical_gaps'), ('다른 가능한 설명', 'alternatives'), ('확인하지 못한 부분', 'gaps'), ('권고 조치', 'actions')]:
        if body.get(field):
            result.append((title, _text_list(body[field])))
    if body.get('timeline'):
        result.append(('핵심 시간축', [cited(f"{x['display_time']} · {x['description']}", x['evidence_ids']) for x in body['timeline']]))
    if audience == 'analyst':
        result.append(('조사 방법', _text_list(body.get('methods')) or ['아래 출처의 도구 입력·결과는 같은 스냅샷의 JSON 근거 부록에 보존되어 있습니다.']))
        result.append(('근거 목록', [f"[{i+1}] {r['data']['tool']} · {r['data'].get('path', '')} {r['data'].get('file_path', '')}\nE:{r['id']}" for i, r in enumerate(body['sources'])]))
    elif body['sources']:
        result.append(('핵심 근거', [f"[{i+1}] {r['data'].get('path', r['data']['tool'])} {r['data'].get('file_path', '')}" for i, r in enumerate(body['sources'])]))
    return result


def render(case, args):
    body, sid = prepare(case, args)
    audience = args.get('audience', 'both')
    if audience not in ('both', 'executive', 'analyst'):
        raise ValueError('audience must be both, executive or analyst')
    outputs = []
    for view in ('executive', 'analyst') if audience == 'both' else (audience,):
        stem = f'report-{sid[:16]}-{view}'
        title = str(body.get('title') or '조사 보고서') + (' 임원 요약' if view == 'executive' else ' 분석가 상세')
        blocks = sections(body, view)
        meta = f"{body['case'].get('case_id', '')} · {body['created_at']} · " + ('부분 조사' if body['analysis_status'] == 'partial' else '요청 범위 조사')
        css = 'body{font:16px/1.65 sans-serif;color:#202329;max-width:850px;margin:36px auto;padding:0 24px;overflow-wrap:anywhere}h1{font-size:28px}h2{font-size:19px;margin-top:26px}small{color:#555}p{white-space:pre-wrap}@media print{body{margin:0;font-size:11pt}h2{break-after:avoid}p{orphans:3;widows:3}}'
        html = f'<!doctype html><html lang="ko"><meta charset="utf-8"><title>{escape(title)}</title><style>{css}</style><body><h1>{escape(title)}</h1><small>{escape(meta)}</small>'
        html += ''.join('<section><h2>' + escape(heading) + '</h2>' + ''.join('<p>' + escape(p) + '</p>' for p in paras if p) + '</section>' for heading, paras in blocks)
        html += f'<p><small>Snapshot {sid}</small></p></body></html>'
        (case.output / (stem + '.html')).write_text(html, encoding='utf-8')
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.72)
        section.left_margin = section.right_margin = Inches(.8)
        for style in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
            s = doc.styles[style]
            s.font.name = 'Noto Sans CJK KR'
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK KR')
            for child in list(s.element.get_or_add_rPr()):
                if child.tag == qn('w:spacing'):
                    s.element.rPr.remove(child)
            ppr = s.element.get_or_add_pPr()
            for child in list(ppr):
                if child.tag == qn('w:pBdr'):
                    ppr.remove(child)
        doc.styles['Normal'].font.size = Pt(11)
        doc.styles['Normal'].paragraph_format.space_after = Pt(5)
        doc.styles['Normal'].paragraph_format.line_spacing = 1.1
        doc.styles['Title'].font.size = Pt(18)
        doc.styles['Title'].paragraph_format.space_after = Pt(8)
        doc.styles['Heading 1'].font.size = Pt(14)
        doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
        doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
        doc.add_paragraph(str(body.get('title') or '조사 보고서'), 'Title')
        doc.add_paragraph('임원용 요약' if view == 'executive' else '분석가용 상세')
        doc.add_paragraph(meta)
        for heading, paras in blocks:
            doc.add_heading(heading, 1)
            for p in paras:
                if p:
                    doc.add_paragraph(p)
        doc.core_properties.title = title
        doc.core_properties.author = 'Forsic'
        doc.core_properties.subject = 'snapshot:' + sid
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        footer.add_run('Forsic · ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)
        doc.save(case.output / (stem + '.docx'))
        outputs.append({'audience': view, 'html': stem + '.html', 'docx': stem + '.docx'})
    return {'snapshot_id': sid, 'source_appendix': 'snapshot-' + sid + '.json', 'reports': outputs,
            'evidence_ids': [r['id'] for r in body['sources']], 'analysis_status': body['analysis_status']}
