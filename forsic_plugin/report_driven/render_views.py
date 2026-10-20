#!/usr/bin/env python3
"""Render the same immutable view through the retained report template.

The DOCX package supplies actual paragraph/table components, page geometry,
styles, headers and footer. HTML mirrors that template's measured layout.
Neither renderer infers facts or performs investigation/model calls.
"""
from __future__ import annotations
import argparse, hashlib, html, io, json, math, re, zipfile
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

STYLE_PATH=Path(__file__).with_name('report_style.json')
STYLE_BYTES=STYLE_PATH.read_bytes()
STYLE=json.loads(STYLE_BYTES)
TEMPLATE_PATH=Path(__file__).parent/STYLE['template_file']
FONT=STYLE['font']['family']


def render_style_identity():
 return {'profile_id':STYLE['profile_id'],'version':STYLE['version'],
         'source_kind':STYLE['source_kind'],'source_file':STYLE_PATH.name,
         'sha256':hashlib.sha256(STYLE_BYTES).hexdigest(),
         'template_file':STYLE['template_file'],
         'template_sha256':hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()}

CSS=f'''
:root{{--ink:#{STYLE['font']['color']};--muted:#{STYLE['text']['muted_color']};--line:#{STYLE['table']['border_color']};--dark:#{STYLE['table']['header_fill']};--wash:#{STYLE['table']['callout_fill']}}}
*{{box-sizing:border-box}}body{{margin:0;background:#ebebeb;color:var(--ink);font-family:"{FONT}","Malgun Gothic",sans-serif;font-size:{STYLE['text']['body_pt']}pt;line-height:{STYLE['text']['line_spacing']}}}
nav{{max-width:210mm;margin:20px auto 10px;display:flex;gap:12px;flex-wrap:wrap;padding:0 12px;font-size:9pt}}a{{color:inherit}}a:focus-visible{{outline:2px solid #555;outline-offset:3px}}
.sheet{{width:210mm;min-height:297mm;margin:0 auto 20px;background:#fff;padding:{STYLE['margins_mm']['top']}mm {STYLE['margins_mm']['right']}mm {STYLE['margins_mm']['bottom']}mm {STYLE['margins_mm']['left']}mm;box-shadow:0 2px 16px #0000000d;display:flex;flex-direction:column}}
.sheet.appendix{{min-height:0}}.running{{text-align:right;font-size:8pt;font-weight:700;color:var(--muted);margin:-9mm 0 6mm}}
.kicker{{font-size:8pt;font-weight:700;color:var(--muted);margin:0 0 4pt}}h1{{color:#{STYLE['text']['title_color']};font-size:{STYLE['text']['title_pt']}pt;font-weight:700;line-height:1;margin:0 0 5pt;padding:0 0 4pt;border-bottom:1pt solid #{STYLE['text']['rule_color']}}}
h2{{color:#{STYLE['text']['title_color']};font-size:{STYLE['text']['heading_pt']}pt;line-height:1.23;margin:10pt 0 5pt;font-weight:700;break-after:avoid}}h3{{color:#{STYLE['text']['title_color']};font-size:{STYLE['text']['subheading_pt']}pt;line-height:1.23;margin:10pt 0 5pt;break-after:avoid}}
p{{margin:0 0 4pt;overflow-wrap:anywhere;white-space:pre-line}}.note{{font-size:8pt;color:var(--muted);margin:0 0 3pt}}
.callout{{background:var(--wash);border:.5pt solid var(--line);padding:4pt 5pt;margin:0 0 12pt;break-inside:auto}}.callout strong{{display:block;font-size:9.5pt;margin-bottom:4pt}}.callout p{{margin:0}}
table{{border-collapse:collapse;width:100%;table-layout:fixed;font-size:{STYLE['table']['body_pt']}pt;line-height:{STYLE['table']['line_spacing']};margin:0 0 12pt}}th,td{{border:.5pt solid var(--line);padding:2pt 5pt;vertical-align:middle;overflow-wrap:anywhere;white-space:pre-line}}th{{font-size:8pt;font-weight:700;background:var(--dark);color:white;text-align:left}}tbody tr:nth-child(even){{background:#{STYLE['table']['stripe_fill']}}}table.metadata td,table.impact td{{font-size:8.5pt}}thead{{display:table-header-group}}tr{{break-inside:avoid}}
caption{{text-align:left;font-size:8pt;color:var(--muted);padding:0 0 3pt}}
.sheet-foot{{margin-top:auto;padding-top:10mm;font-size:8pt;color:var(--muted);text-align:right}}.page-body{{min-width:0}}
@page{{size:A4;margin:{STYLE['margins_mm']['top']}mm {STYLE['margins_mm']['right']}mm {STYLE['margins_mm']['bottom']}mm {STYLE['margins_mm']['left']}mm}}
@media print{{body{{background:#fff}}nav{{display:none}}.sheet{{width:auto;min-height:0;display:block;margin:0;padding:0;box-shadow:none;break-after:page}}.sheet:last-child,.sheet.appendix{{break-after:auto}}.sheet.appendix+.sheet.appendix{{padding-top:5mm}}.running{{margin:0 0 6mm}}.sheet-foot{{padding-top:7mm}}a{{text-decoration:none}}}}
@media screen and (max-width:850px){{body{{font-size:15px}}.sheet{{width:100%;min-height:0;padding:24px 20px;margin-bottom:12px}}nav{{margin:12px}}.running{{margin:0 0 18px}}h1{{font-size:25px}}h2{{font-size:20px}}h3{{font-size:17px}}table,table.metadata td,table.impact td,th{{font-size:13px}}th,td{{padding:7px 6px}}.note,.kicker,.sheet-foot,.running{{font-size:12px}}.callout strong{{font-size:15px}}.page-body{{overflow-x:auto}}}}
@media(prefers-reduced-motion:reduce){{*{{scroll-behavior:auto!important;animation:none!important}}}}
'''


def text(v):return html.escape(str(v),quote=True)
def bookmark_id(v):return 'r_'+hashlib.sha256(str(v).encode()).hexdigest()[:28]
def linked_text(value,targets):
 value=str(value)
 if not targets:return text(value)
 pattern='('+'|'.join(re.escape(x) for x in sorted(targets,key=len,reverse=True))+')'
 return ''.join('<a href="#'+bookmark_id(part)+'">'+text(part)+'</a>' if part in targets else text(part) for part in re.split(pattern,value))


def hblock(b,targets=(),anchors=None):
 if b.get('literal'):targets=()
 kind=b['kind']
 if kind=='heading':
  tag='h3' if b.get('level')==2 else 'h2'
  ident=' id="'+bookmark_id(b['id'])+'"' if b.get('id') else ''
  return '<'+tag+ident+'>'+text(b['text'])+'</'+tag+'>'
 if kind in ('paragraph','note'):return '<p'+(' class="note"' if kind=='note' else '')+(' id="'+bookmark_id(b['id'])+'"' if b.get('id') else '')+'>'+linked_text(b['text'],targets)+'</p>'
 if kind=='callout':return '<aside class="callout"><strong>'+text(b['label'])+'</strong><p>'+linked_text(b['text'],targets)+'</p></aside>'
 if kind=='table':
  widths=b.get('widths') or [100/len(b['headers'])]*len(b['headers'])
  ids=b.get('row_ids',[]);rows=[]
  for i,row in enumerate(b['rows']):
   ident=''
   if i<len(ids) and (anchors is None or ids[i] not in anchors):
    ident=' id="'+bookmark_id(ids[i])+'"'
    if anchors is not None:anchors.add(ids[i])
   rows.append('<tr'+ident+'>'+''.join('<td>'+linked_text(x,targets)+'</td>' for x in row)+'</tr>')
  variant=b.get('variant','detail')
  caption=('<caption id="'+bookmark_id(b['id'])+'">'+text(b['id'])+'</caption>') if b.get('id') else ''
  return '<table class="'+text(variant)+'">'+caption+'<colgroup>'+''.join(f'<col style="width:{float(w)}%">' for w in widths)+'</colgroup><thead><tr>'+''.join('<th scope="col">'+text(x)+'</th>' for x in b['headers'])+'</tr></thead><tbody>'+''.join(rows)+'</tbody></table>'
 raise ValueError('Unsupported block '+str(kind))


def report_status(v):
 return {'template':'미기입 표준양식','synthetic':'가상자료 예시','live':'판단 기준시점의 보고서'}[v['data_mode']]


def render_html(v):
 targets={x for p in v['pages'] for b in p['blocks'] for x in [*b.get('row_ids',[]),*([b['id']] if b.get('id') else [])]}
 nav='<nav aria-label="보고서 섹션">'+''.join('<a href="#'+bookmark_id(p['id'])+'">'+text(p['title'])+'</a>' for p in v['pages'])+'</nav>'
 body=[];anchors=set()
 for i,p in enumerate(v['pages'],1):
  kicker=p.get('kicker') or ('통합보고서 / 실무 상세 부록' if p.get('section')=='appendix' else '통합보고서')
  body.append('<article class="sheet'+(' appendix' if p.get('section')=='appendix' else '')+'" id="'+bookmark_id(p['id'])+'"><header class="running">FORSIC | 침해사고 조사 통합보고서 | '+text(report_status(v))+'</header><div class="page-body"><p class="kicker">'+text(kicker)+'</p><h1>'+text(p['title'])+'</h1>'+''.join(hblock(b,targets,anchors) for b in p['blocks'])+'</div><footer class="sheet-foot">'+text(v['security']+' | '+report_status(v))+' · '+str(i)+'</footer></article>')
 return '<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; img-src \'none\'; base-uri \'none\'; form-action \'none\'"><title>'+text(v['title'])+'</title><style>'+CSS+'</style></head><body>'+nav+''.join(body)+'</body></html>'


def _fill_paragraph(p,value,targets=()):
 """Replace the editable text, retaining the source paragraph/run formatting."""
 runs=list(p.runs)
 props=next((r._r.rPr for r in runs if r.text and r._r.rPr is not None),None)
 if props is None:props=next((r._r.rPr for r in runs if r._r.rPr is not None),None)
 props=deepcopy(props)
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 pattern='('+'|'.join(re.escape(x) for x in sorted(targets,key=len,reverse=True))+')' if targets else None
 for part in re.split(pattern,str(value)) if pattern else [str(value)]:
  run=p.add_run(part)
  if props is not None:run._r.insert(0,deepcopy(props))
  if part in targets:
   link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),bookmark_id(part));p._p.remove(run._r);link.append(run._r);p._p.append(link)
 return p


def _bookmark(p,key,seen):
 if key in seen:return
 seen.add(key);ident=str(len(seen))
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),ident);start.set(qn('w:name'),bookmark_id(key))
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),ident)
 p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end)


def _cell_text(cell,value,targets):
 ps=list(cell.paragraphs)
 for p in ps[1:]:cell._tc.remove(p._p)
 return _fill_paragraph(ps[0],value,targets)


def _table(doc,prototype,b,targets,seen,style):
 headers=b['headers'];n=len(headers)
 if not n or any(len(row)!=n for row in b['rows']):raise ValueError('Table column mismatch')
 element=deepcopy(prototype);table=Table(element,doc._body)
 samples=[deepcopy(row._tr) for row in table.rows]
 old_n=len(table.columns)
 for row in list(table.rows):element.remove(row._tr)
 widths=b.get('widths') or [100/n]*n
 if len(widths)!=n or any(w<=0 for w in widths):raise ValueError('Invalid table widths')
 fixed=bool(b.get('variant')) and old_n==n
 if not fixed:
  grid=element.tblGrid
  for child in list(grid):grid.remove(child)
  for width in widths:
   col=OxmlElement('w:gridCol');col.set(qn('w:w'),str(round(9752*width/sum(widths))));grid.append(col)
 line_estimates=[max(sum(max(1,math.ceil(len(line)/max(1,(style['table']['width_mm']*w/sum(widths)-4)*2.835/style['table']['body_pt']))) for line in str(v).split('\n')) for v,w in zip(values,widths)) for values in [headers,*b['rows']]]
 keep_together=b.get('keep_together') and sum(line_estimates)<35
 for ri,values in enumerate([headers,*b['rows']]):
  source=samples[0] if ri==0 else samples[min(1+(ri-1)%2,len(samples)-1)]
  row=deepcopy(source)
  cells=list(row.tc_lst)
  while len(cells)<n:row.append(deepcopy(cells[-1]));cells=list(row.tc_lst)
  for cell in cells[n:]:row.remove(cell)
  element.append(row)
  trpr=row.get_or_add_trPr()
  for flag in list(trpr):
   if flag.tag in (qn('w:tblHeader'),qn('w:cantSplit')):trpr.remove(flag)
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  # Preserve ordinary row integrity, but allow long exact original text to flow.
  longest=line_estimates[ri]
  if ri==0 or longest<30:trpr.append(OxmlElement('w:cantSplit'))
  for ci,(cell,value,width) in enumerate(zip(table.rows[-1].cells,values,widths)):
   pr=cell._tc.get_or_add_tcPr()
   if not fixed:pr.find(qn('w:tcW')).set(qn('w:w'),str(round(9752*width/sum(widths))))
   for shade in list(pr.findall(qn('w:shd'))):pr.remove(shade)
   fill=style['table']['header_fill'] if ri==0 else style['table']['stripe_fill'] if ri%2==0 else None
   if fill:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);pr.append(sh)
   para=_cell_text(cell,value,targets if ri else ())
   if keep_together:para.paragraph_format.keep_with_next=ri<len(b['rows'])
   if ri and ci==0 and ri<=len(b.get('row_ids',[])):_bookmark(para,b['row_ids'][ri-1],seen)
 doc.element.body.insert(-1,element)
 return table


def render_docx(v,path,*,template_bytes=None,style_profile=None):
 # Missing or corrupt template is an error, never a silent generic fallback.
 original=TEMPLATE_PATH.read_bytes() if template_bytes is None else template_bytes
 style=deepcopy(STYLE if style_profile is None else style_profile)
 doc=Document(io.BytesIO(original));body=doc.element.body
 components={name:deepcopy(body[index]) for name,index in style['components'].items()}
 for child in list(body):
  if child.tag!=qn('w:sectPr'):body.remove(child)
 targets={x for p in v['pages'] for b in p['blocks'] for x in [*b.get('row_ids',[]),*([b['id']] if b.get('id') else [])]};seen=set()
 def paragraph(role,value,links=True):
  el=deepcopy(components[role]);body.insert(-1,el)
  return _fill_paragraph(Paragraph(el,doc._body),value,targets if links else ())
 def spacer():body.insert(-1,deepcopy(components['spacer']))
 # Only the template's content/status fields change; placement and PAGE field stay.
 for section in doc.sections:
  for p in section.header.paragraphs:
   if 'FORSIC' in p.text:_fill_paragraph(p,'FORSIC  |  침해사고 조사 통합보고서  |  '+report_status(v))
  for p in section.footer.paragraphs:
   for node in p._p.xpath('.//w:t'):
    if '미기입 표준양식' in (node.text or ''):node.text=v['security']+'  |  '+report_status(v)+'  ·  '
 cp=doc.core_properties;cp.title=v['title'];cp.subject='FORSIC snapshot:'+v['snapshot_revision']
 cp.author='';cp.last_modified_by='';cp.comments=''
 for i,page in enumerate(v['pages']):
  continuing=i>0 and page.get('section')=='appendix' and v['pages'][i-1].get('section')=='appendix'
  if not continuing:
   p=paragraph('kicker',page.get('kicker') or ('통합보고서 / 실무 상세 부록' if page.get('section')=='appendix' else '통합보고서'),False)
   p.paragraph_format.page_break_before=bool(i);p.paragraph_format.keep_with_next=True
   p=paragraph('title',page['title'],False)
  else:
   p=paragraph('heading',page['title'],False)
  _bookmark(p,page['id'],seen)
  for block in page['blocks']:
   kind=block['kind']
   if kind in ('paragraph','note'):
    p=paragraph(kind,block['text'])
    if block.get('id'):_bookmark(p,block['id'],seen)
   elif kind=='heading':
    p=paragraph('subheading' if block.get('level')==2 else 'heading',block['text'])
    if block.get('id'):_bookmark(p,block['id'],seen)
   elif kind=='table':
    if block.get('id'):_bookmark(paragraph('note',block['id'],False),block['id'],seen)
    variant=block.get('variant');component=variant if variant in components else 'flow' if len(block['headers'])==3 else 'summary'
    _table(doc,components[component],block,() if block.get('literal') else targets,seen,style);spacer()
   elif kind=='callout':
    el=deepcopy(components['callout']);body.insert(-1,el);table=Table(el,doc._body)
    ps=table.cell(0,0).paragraphs
    _fill_paragraph(ps[0],block['label']);_fill_paragraph(ps[1],block['text'],targets)
    for p in ps[2:]:p._p.getparent().remove(p._p)
    spacer()
   else:raise ValueError('Unsupported block '+str(kind))
 generated=io.BytesIO();doc.save(generated)
 # python-docx writes only supported parts back through its object model. Keep
 # every unrelated original package part byte-identical, including the real
 # template styles/theme/numbering and opaque XML. No external relationships.
 changed={'word/document.xml','word/header1.xml','word/footer1.xml','docProps/core.xml'}
 with zipfile.ZipFile(io.BytesIO(original)) as source,zipfile.ZipFile(generated) as output,zipfile.ZipFile(path,'w') as final:
  if set(source.namelist())!=set(output.namelist()):raise ValueError('Unexpected template package parts')
  for item in source.infolist():
   payload=output.read(item.filename) if item.filename in changed else source.read(item.filename)
   if item.filename.endswith('.rels') and b'TargetMode="External"' in payload:raise ValueError('External relationship')
   final.writestr(item,payload)


def render_pair(view_file,outdir):
 v=json.loads(Path(view_file).read_text(encoding='utf-8'));out=Path(outdir);out.mkdir(parents=True,exist_ok=True)
 stem=Path(view_file).name.removesuffix('.view.json')+'_template'
 hp,dp=out/(stem+'.html'),out/(stem+'.docx')
 if hp.exists() or dp.exists():raise FileExistsError('Create-only outputs: choose a new directory')
 hp.write_text(render_html(v),encoding='utf-8');render_docx(v,dp)
 return hp,dp

if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('view');ap.add_argument('--out',required=True);args=ap.parse_args()
 for p in render_pair(args.view,args.out):print(p)
