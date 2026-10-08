"""Offline host acceptance: no model, E01, service restart or network."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from forsic_plugin.evidence import Case
from forsic_plugin.report_driven.host import current, invoke, ref, checked, ledger
from forsic_plugin.report_driven.views import bundle, compile_views, bundles, artifact
from forsic_plugin.report_driven.gap_audit import audit
from forsic_plugin.notes import note


def fixture(root):
    root=Path(root);e=root/'evidence';e.mkdir(parents=True)
    (e/'request.log').write_text('job=maintenance requested target=/opt/task\n')
    (e/'process.log').write_text('process=task run_id=R-77 state=created\n')
    manifest=root/'case.json';manifest.write_text(json.dumps(dict(case_id='SYN-report-driven',label='가공 작업 기록',evidence_root=str(e),output_root=str(root/'results'),scope='가공 요청·프로세스 두 레코드의 연결',question='요청 기록과 프로세스 생성 기록은 같은 작업인가?',synthetic=True)))
    c=Case(manifest)
    records=[json.loads(c.invoke('forsic_read',dict(path=p))) for p in ('request.log','process.log')]
    n=note(c,dict(action='save',question='요청과 생성 기록이 같은 작업인가?',answer='이름은 비슷하지만 공통 작업 식별이 아직 확인되지 않았습니다.',evidence_ids=[r['evidence_id'] for r in records],alternatives=['서로 다른 작업의 기록일 수 있습니다.'],critical_gaps=['공통 작업 identity 연결 미확인'],gaps=['승인·최종 실행 성공은 범위 밖'],next_checks=['공통 작업 식별을 포함한 기록 확인']))['note']
    return c,records,n


def mission(s):
    q=s['questions'][0];g=s['gaps'][0];src=s['sources'][-1]
    return dict(id='M-link',version='1',original_obligation_id=g['id'],gap_ref=ref('gap',g),question_ref=ref('question',q),report_targets=g['section_targets'],target_proposition=q['question'],why_it_matters='설정·요청과 실제 생성을 구별하기 위해 필요합니다.',current_answer=q['answer'],competing_explanation_refs=q['hypothesis_refs'],mission_kind='evaluate_result',input_refs=[ref('source',r) for r in s['sources']],required_view='full_field',capability_requirement='보존된 두 레코드에서 공통 작업 식별 비교',proposed_tool_name=None,exact_target_scope=q['scope'],reuse_result_refs=[ref('source',src)],design_timing='after_result',support_rule='두 기록의 공통 작업 식별과 대상이 일치',refute_rule='작업 식별이 명시적으로 다름',inconclusive_rule='이름만 같거나 공통 식별이 없으면 판별 불가',preserved_counterevidence_refs=[ref('source',s['sources'][0])],state='draft',budget=dict(authority='unallocated'),material_change_required='기존 반환에서 연결 조건 평가; 물리 재실행 없음',completion_proof=['정확 결과판 평가와 질문 답 갱신'],does_not_resolve=['승인 여부','최종 실행 성공','악성 목적'],reopen_conditions=['새 공통 작업 식별이 포함된 원문'])


def proposal(s,c):
    m=s['missions'][0];src=s['sources'][-1]
    return dict(mission_id=m['id'],mission_version=m['version'],result_id=src['id'],result_version=src['version'],outcome='inconclusive',reasoning_summary='프로세스 기록에는 run_id가 있으나 요청 기록에는 없어 동일 작업인지 연결할 수 없습니다.',answer='요청과 프로세스 생성 기록은 있으나 두 기록의 동일 작업 여부는 판별하지 못했습니다.',citations=[dict(source_id=x['id'],source_version=x['version'],pointer='/lines/0/text',literal=c.event(x['id'])['data']['lines'][0]['text'],byte_start=0) for x in s['sources']],limitations=['공통 작업 식별이 없어 이름의 유사성만으로 연결할 수 없습니다.'],next_check='공통 run_id가 포함된 요청 이력을 얻으면 다시 판별합니다.')


def worked(c):
    s=current(c);invoke(c,dict(action='mission',snapshot_id=s['meta']['snapshot_id'],payload=mission(s)))
    s=current(c);invoke(c,dict(action='assess',snapshot_id=s['meta']['snapshot_id'],payload=proposal(s,c)))
    return current(c)


class ReportDrivenTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.c,self.results,self.n=fixture(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def call(self,action,p=None):
        return invoke(self.c,dict(action=action,snapshot_id=current(self.c)['meta']['snapshot_id'],payload=p if p is not None else {}))
    def start(self):self.call('mission',mission(current(self.c)));return current(self.c)
    def test_T01_empty_not_final(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);(root/'e').mkdir();p=root/'c.json';p.write_text(json.dumps(dict(case_id='empty',evidence_root=str(root/'e'),output_root=str(root/'o'),synthetic=True)))
            a=audit(current(Case(p)));self.assertTrue(a['scope_final_blockers']);self.assertEqual(a['publication_recommendation'],'draft_only')
    def test_missing_goal_is_unknown_not_stopped(self):
        self.assertEqual(current(self.c)['status']['execution'],'unknown')
    def test_T02_unknown_needs_reason_scope_reopen(self):
        s=self.start();p=proposal(s,self.c);p['next_check']=''
        with self.assertRaises(ValueError):self.call('assess',p)
        p=proposal(s,self.c);self.call('assess',p);s=current(self.c)
        self.assertEqual(s['questions'][0]['assessment'],'undetermined');self.assertTrue(s['gaps'][0]['reopen_conditions'])
    def test_T03_returned_not_closed(self):
        s=current(self.c);self.assertTrue(audit(s)['ready_to_plan_gap_ids']);self.assertFalse(s['assessments']);self.assertEqual(s['questions'][0]['work_state'],'open')
    def test_T04_discovery_not_support(self):
        s=current(self.c);m=mission(s);m['support_rule']=m['refute_rule']=None;self.call('mission',m);s=current(self.c);p=proposal(s,self.c);p['outcome']='supports'
        with self.assertRaises(ValueError):self.call('assess',p)
    def test_T05_exact_body_not_metadata_or_bad_span(self):
        s=self.start();p=proposal(s,self.c);p['citations'][0].update(pointer='/path',literal='request.log')
        with self.assertRaises(ValueError):self.call('assess',p)
        p=proposal(s,self.c);p['citations'][0]['literal']='invented'
        with self.assertRaises(ValueError):self.call('assess',p)
    def test_T06_internal_error_not_external(self):
        g=deepcopy(current(self.c)['gaps'][0]);g.update(kind='internal_error',feasible_next_action='request_external_input')
        with self.assertRaises(ValueError):self.call('gap',g)
        g['feasible_next_action']='repair_internal';self.call('gap',g)
        self.assertNotIn(g['id'],self.call('gaps')['audit']['external_input_gap_ids'])
    def test_T07_same_obligation_survives_assessment(self):
        s=worked(self.c);self.assertFalse(s['assessments'][0]['resolved_gap_ids']);g=next(x for x in s['gaps'] if x['id']==s['missions'][0]['gap_ref']['id']);self.assertEqual(g['disposition'],'assessed_unresolved');self.assertEqual(s['missions'][0]['state'],'partial')
    def test_T08_counterevidence_every_reader(self):
        s=worked(self.c)
        for v in compile_views(s).values():
            text=json.dumps(v,ensure_ascii=False);self.assertIn('서로 다른 작업',text);self.assertIn('공통 작업 식별',text)
    def test_T09_reuse_not_new_physical_result(self):
        s=self.start();p=proposal(s,self.c);self.call('assess',p);p['mission_version']=current(self.c)['missions'][0]['version'];again=self.call('assess',p)
        self.assertTrue(again['reused']);self.assertEqual(len(current(self.c)['sources']),2);self.assertEqual(len(current(self.c)['assessments']),1)
    def test_T10_unknown_time_never_normalized(self):
        note(self.c,dict(action='save',note_id=self.n['note_id'],revision=1,question=self.n['question'],answer=self.n['answer'],evidence_ids=[self.results[0]['evidence_id']],timeline=[dict(time='Oct 1 03:00',description='시간대 미상',evidence_ids=[self.results[0]['evidence_id']])]))
        s=current(self.c);self.assertFalse(s['timeline']);self.assertIn('Oct 1',str(s['status']['explicit_limitations']))
    def test_T11_recommended_not_performed(self):self.assertEqual(current(self.c)['actions'][0]['action_kind'],'recommended')
    def test_T12_four_outputs_same_content_and_manifest(self):
        from docx import Document
        s=worked(self.c);r=bundle(self.c,s);m=json.loads(Path(r['manifest']).read_text());self.assertEqual(len(m['files']),4)
        for name in r['files']:
            p=Path(name);text=p.read_text() if p.suffix=='.html' else '\n'.join(x.text for x in Document(p).paragraphs)+str([[c.text for c in row.cells] for t in Document(p).tables for row in t.rows])
            self.assertIn('F-M-link',text);self.assertIn('동일 작업 여부는 판별하지 못했습니다',text);self.assertIn('공통 작업 식별',text)
        self.assertEqual(r['snapshot_id'],s['meta']['snapshot_id'])
    def test_T14_correction_preserves_old_and_stales_current(self):
        s=worked(self.c);r=bundle(self.c,s);before={p:Path(p).read_bytes() for p in r['files']}
        latest=note(self.c,dict(action='get',note_id=self.n['note_id']))['note']
        note(self.c,dict(action='save',note_id=self.n['note_id'],revision=latest['revision'],question='정정된 질문',answer='미확인',evidence_ids=[self.results[0]['evidence_id']],correction_reason='정의 변경'))
        now=current(self.c);self.assertFalse(now['assessments']);self.assertTrue(bundles(self.c)[0]['stale']);checked(now)
        self.assertEqual(before,{p:Path(p).read_bytes() for p in before})
    def test_T15_escape_and_download_path(self):
        from forsic_plugin.report_driven.render_views import render_html
        s=current(self.c);s['claims'][0]['statement']='<script src="https://bad.invalid">attack</script>'
        text=render_html(compile_views(s)['executive']);self.assertNotIn('<script',text);self.assertIn('&lt;script',text)
        with self.assertRaises(ValueError):artifact(self.c,'../','manifest.json')
    def test_T16_passive_reads_no_events_or_model(self):
        before=self.c.db.read_bytes()
        with patch('socket.socket',side_effect=AssertionError('network forbidden')):
            for _ in range(2):self.call('state');self.call('gaps');bundles(self.c)
        self.assertEqual(before,self.c.db.read_bytes())
    def test_T17_and_T18_no_status_authority(self):
        self.assertFalse(self.call('publish')['published']);s=self.start();p=proposal(s,self.c);p['adoption_receipt']='fake'
        with self.assertRaises(ValueError):self.call('assess',p)
        self.assertEqual(s['meta']['data_mode'],'synthetic');self.assertEqual(s['status']['approval_gate'],'not_requested')
    def test_stale_snapshot_and_foreign_refs(self):
        s=current(self.c);m=mission(s);m['reuse_result_refs'][0]['version']='wrong'
        with self.assertRaises(ValueError):self.call('mission',m)
        with self.assertRaises(ValueError):invoke(self.c,dict(action='mission',snapshot_id='old',payload=mission(s)))
    def test_no_self_assigned_budget(self):
        m=mission(current(self.c));m['state']='ready';m['budget']=dict(authority='host_reserved',model_calls=2,input_tokens=1000,output_tokens=1000,wall_seconds=60)
        with self.assertRaises(ValueError):self.call('mission',m)
    def test_mask_copy_and_immutable_bundle_reuse(self):
        s=worked(self.c);before=deepcopy(s);r=bundle(self.c,s,['maintenance']);self.assertEqual(before,s)
        self.assertEqual(r,bundle(self.c,s,['maintenance']))
        self.assertIn('[비공개]',Path(r['files'][2]).read_text())
    def test_unassessed_conditional_not_na(self):
        r=current(self.c)['requirements'];self.assertTrue(any(x['applicability']=='conditional_unassessed' for x in r));self.assertFalse(any(x['applicability']=='not_applicable' for x in r))
    def test_material_requirement_unknown_needs_gaps(self):
        r=deepcopy(current(self.c)['requirements'][0]);r.update(disposition='scoped_unknown',gap_ids=[])
        with self.assertRaises(ValueError):self.call('requirement',r)
    def test_skill_selected_questions_and_no_business_recommendations(self):
        note(self.c,dict(action='save',question='개인정보 보유 여부는?',answer='내용을 아직 확인하지 못했습니다.',evidence_ids=[self.results[0]['evidence_id']],critical_gaps=['본문의 정보 종류 미확인'],next_checks=['대상 본문 종류 확인']))
        s=current(self.c);q=next(q for q in s['questions'] if q['question']=='개인정보 보유 여부는?')
        r=deepcopy(next(r for r in s['requirements'] if r['requirement_id']=='REQ-CORE-ANSWER'))
        r.update(basis_refs=[ref('question',q)],rationale='사용자가 확인하려는 정보 보유 범위',gap_ids=q['remaining_gap_ids'])
        self.call('requirement',r);s=current(self.c)
        s['actions'].append(dict(action_kind='recommended',title='경영 권고 고유문구',basis_refs=[],execution_receipt=None))
        views=compile_views(s);rows=views['executive']['pages'][0]['blocks'][0]['rows']
        self.assertEqual([row[0] for row in rows],['개인정보 보유 여부는?'])
        self.assertNotIn('요청과 생성 기록이 같은 작업인가?',json.dumps(views['executive']['pages'][0],ensure_ascii=False))
        for v in views.values():self.assertNotIn('경영 권고 고유문구',json.dumps(v,ensure_ascii=False))
        self.assertNotIn('의사결정',json.dumps(views['executive'],ensure_ascii=False))
        self.assertTrue(q['remaining_gap_ids']);self.assertIn('본문의 정보 종류 미확인',json.dumps(views,ensure_ascii=False))
    def test_actual_native_tool_path(self):
        r=json.loads(self.c.invoke('forsic_reporting',{'action':'state'}));self.assertNotIn('error',r);self.assertIn('state',r)
    def test_native_manifest_and_schema_references(self):
        from unittest.mock import Mock
        from ruamel.yaml import YAML
        from forsic_plugin import register
        ctx=Mock();ctx.get_config.side_effect=lambda k,d=None: str(Path(self.tmp.name)/'intake') if k=='intake_root' else d
        register(ctx)
        registrations={c.kwargs['name']:c.kwargs['schema'] for c in ctx.register_tool.call_args_list}
        manifest=YAML(typ='safe').load((Path(__file__).parent/'forsic_plugin/plugin.yaml').read_text())
        self.assertEqual(set(manifest['provides_tools']),set(registrations))
        schema=registrations['forsic_reporting']['parameters']
        def check(value):
            if isinstance(value,dict):
                if '$ref' in value:self.assertIn(value['$ref'].removeprefix('#/$defs/'),schema['$defs'])
                for v in value.values():check(v)
            elif isinstance(value,list):
                for v in value:check(v)
        check(schema)
    def test_business_question_gap_mission_assessment_loop(self):
        note(self.c,dict(action='save',question='개인정보가 외부로 전송됐는가?',answer='제시된 두 작업 기록으로는 알 수 없습니다.',evidence_ids=[self.results[0]['evidence_id']],critical_gaps=['전송 내용과 대상 연결 미확인'],next_checks=['이미 보존된 기록의 전송 정보 판별']))
        s=current(self.c);q=next(q for q in s['questions'] if q['question']=='개인정보가 외부로 전송됐는가?');g=next(g for g in s['gaps'] if g['question_ref']['id']==q['id'])
        m=mission(s);m.update(id='M-transmission',original_obligation_id=g['id'],gap_ref=ref('gap',g),question_ref=ref('question',q),report_targets=g['section_targets'],target_proposition=q['question'],current_answer=q['answer'],competing_explanation_refs=q['hypothesis_refs'],exact_target_scope=q['scope'],support_rule='기록에 개인정보 내용과 송신 성공이 함께 확인됨',refute_rule='해당 송신 내용이 명시적으로 개인정보가 아님',inconclusive_rule='내용·대상·송신 성공이 미제시이면 미확인')
        self.call('mission',m);s=current(self.c);p=proposal(s,self.c)
        p.update(answer='제시된 작업 기록에는 전송 내용이 없어 개인정보 전송 여부를 확인하지 못했습니다.',reasoning_summary='작업 생성 기록은 송신 내용이나 성공을 보여주지 않습니다.',next_check='전송 대상·내용을 판별할 기록을 확보하면 재검토합니다.')
        self.call('assess',p);s=current(self.c);q=next(x for x in s['questions'] if x['id']==q['id'])
        self.assertEqual(q['answer'],p['answer']);self.assertEqual(q['assessment'],'undetermined');self.assertIn(g['id'],q['remaining_gap_ids']);self.assertEqual(len(s['sources']),2)
    def test_T13_links_resolve_and_word_relationships_local(self):
        from forsic_plugin.report_driven.render_views import render_html,bookmark_id
        import re,zipfile
        from xml.etree import ElementTree as ET
        s=worked(self.c);r=bundle(self.c,s)
        html=Path(r['files'][2]).read_text();ids=set(re.findall(r' id="([^"]+)"',html));links=set(re.findall(r'href="#([^"]+)"',html))
        self.assertTrue(links<=ids);self.assertIn(bookmark_id(s['sources'][0]['id']),links)
        with zipfile.ZipFile(r['files'][3]) as z:
            ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};x=ET.fromstring(z.read('word/document.xml'))
            anchors={a.attrib['{'+ns['w']+'}anchor'] for a in x.findall('.//w:hyperlink',ns)}
            names={a.attrib['{'+ns['w']+'}name'] for a in x.findall('.//w:bookmarkStart',ns)}
            self.assertTrue(anchors);self.assertTrue(anchors<=names)
            self.assertFalse(any(b'TargetMode="External"' in z.read(n) for n in z.namelist() if n.endswith('.rels')))
    def test_large_inventory_not_truncated(self):
        for i in range(205):self.c.record('tool_result',{'tool':'forsic_read','started_id':'synthetic-job-'+str(i),'path':'가공/'+('long-path-'*30)+str(i),'lines':[{'line':1,'text':'가공 기록 '+str(i)}]})
        s=current(self.c);v=compile_views(s)['practitioner'];rows=v['pages'][1]['blocks'][-1]['rows'];self.assertEqual(len(rows),207)
        self.assertIn(s['sources'][-1]['id'],json.dumps(v))
    def test_answer_revision_changes_and_old_model_status_not_authority(self):
        s=self.start();old=s['questions'][0]['version'];self.call('assess',proposal(s,self.c));s=current(self.c)
        self.assertNotEqual(old,s['questions'][0]['version']);checked(s)
    def test_http_passive_state_download_and_no_mutation(self):
        import asyncio,httpx
        from fastapi import FastAPI
        from forsic_plugin.dashboard import plugin_api
        s=worked(self.c);result=bundle(self.c,s);before=self.c.db.read_bytes();app=FastAPI();app.include_router(plugin_api.router)
        async def run():
            with patch.object(plugin_api,'selected_case',return_value=self.c):
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app),base_url='http://test') as client:
                    response=await client.get('/report-driven');self.assertEqual(response.status_code,200)
                    self.assertEqual(response.json()['state']['meta']['snapshot_id'],s['meta']['snapshot_id'])
                    for name in ('executive.html','executive.docx','practitioner.html','practitioner.docx','manifest.json'):
                        self.assertEqual((await client.get('/report-bundles/'+result['bundle_id']+'/'+name)).status_code,200)
                    self.assertEqual((await client.get('/report-bundles/'+result['bundle_id']+'/credentials.json')).status_code,404)
        asyncio.run(run());self.assertEqual(before,self.c.db.read_bytes())


if __name__=='__main__':unittest.main()
