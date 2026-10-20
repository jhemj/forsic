"""Offline reader projection and actual style binding contracts."""
from copy import deepcopy
import hashlib
from html import unescape
import json
from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZipFile
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from forsic_plugin.report_driven.host import current
from forsic_plugin.report_driven.reader_projection import labels, reader_text
from forsic_plugin.report_driven.render_views import bookmark_id, render_html, render_docx, render_style_identity
from forsic_plugin.report_driven.views import artifact, bundle, compile_views
from test_report_driven import fixture, worked


def visible_blocks(pages):
    return '\n'.join(str(value) for page in pages for block in page['blocks']
                     for key in ('text','label','headers','rows') if key in block
                     for value in (block[key] if isinstance(block[key],list) else [block[key]]))


def granular_fixture(root):
    case, _, _ = fixture(root)
    state = worked(case)
    claim = state['claims'][-1]
    claim.update(statement='요청 이력은 확인되었으나 실행 성공은 미확인.', assertion_kind='interpretation',
                 inference_strength='favored', limitations=['전체 기간의 성공 여부는 미확인.'],
                 assumptions=['기록에 제시된 대상 범위만 적용.'],
                 declared_analysis_fields=['counter_citation_refs','inference_strength','assumptions','limitations'])
    return case, state


def declared_view(state, *, needs_recheck=False, declared_state="inferred"):
    # Explicit analyst declaration is separate from ordinary legacy prose. This
    # renderer test supplies the public projection API; writer tests validate it.
    keys=('compromise_time','entry_route','attacker_actions','personal_data_presence',
          'external_disclosure','decisive_evidence','incident_timeline')
    claim=state['claims'][-1]
    entry=dict(question_id=state['questions'][0]['id'],state=declared_state,
        brief=claim['statement'],scope=claim['scope'],basis=claim['reasoning_summary'],
        limitations=claim['limitations'],assumptions=claim['assumptions'],
        evidence_refs=claim['source_refs'],claim_refs=[],counterevidence_refs=claim['counterevidence_refs'],needs_recheck=needs_recheck)
    projection={'schema':'forsic-report-slots-1','snapshot_id':state['meta']['snapshot_id'],
                'slots':[dict(slot=key,label=key,missing=key!='attacker_actions',
                              entries=[entry] if key=='attacker_actions' else []) for key in keys]}
    with patch('forsic_plugin.report_driven.report_slots.project_report_slots',return_value=projection):
        return compile_views(state)['integrated']


class ReaderProjectionTests(unittest.TestCase):
    def test_reader_aliases_replace_complete_identities_only(self):
        mapping={'scope':'조사 범위','F-note':'판단 1'}
        self.assertEqual(reader_text('microscope scope_limited F-noteworthy F-note-other',mapping),
                         'microscope scope_limited F-noteworthy F-note-other')
        self.assertEqual(reader_text('(scope), F-note / exact reference',mapping),
                         '(조사 범위), 판단 1 / exact reference')

    def test_legacy_answer_is_preserved_in_appendix_without_classification(self):
        with tempfile.TemporaryDirectory() as root:
            case,_,_=fixture(root); state=current(case)
            answer='긴 답의 전반부.\n'+('세부 설명과 조건을 그대로 보존. '*90)+'\n끝부분의 중요한 유보.'
            state['questions'][0]['answer']=answer
            state['claims'][0]['statement']=answer
            before=deepcopy(state); view=compile_views(state)['integrated']
            body=visible_blocks([p for p in view['pages'] if p['section']=='body'])
            self.assertNotIn(answer,body)
            self.assertIn('보고용 판단 선언',body)
            # The exact canonical answer now appears once in its technical
            # table; inspect rendered text rather than Python's escaped row repr.
            self.assertIn(answer,unescape(re.sub('<[^>]+>','',render_html(view))))
            self.assertEqual(state,before)

    def test_granular_claims_qualifiers_counterevidence_and_scope_remain_visible(self):
        with tempfile.TemporaryDirectory() as root:
            _,state=granular_fixture(root); before=deepcopy(state)
            view=declared_view(state); main=[p for p in view['pages'] if p['section']=='body']
            body=visible_blocks(main)
            for phrase in ('요청 이력은 확인되었으나 실행 성공은 미확인.','전체 기간의 성공 여부는 미확인.',
                           '기록에 제시된 대상 범위만 적용.','추정',state['questions'][0]['scope']):
                self.assertIn(phrase,body)
            self.assertNotIn(state['questions'][0]['answer'],body)
            self.assertIn('핵심 확인내용',body)
            for collection in ('questions','claims','observations','sources','missions','gaps'):
                for row in state[collection]:self.assertNotIn(row['id'],body)
            self.assertEqual(state,before)
            self.assertEqual(len(main),3)
            self.assertIn(state['questions'][0]['answer'],visible_blocks(view['pages'][3:]))

    def test_stale_confirmed_declaration_is_never_labeled_current_confirmed(self):
        with tempfile.TemporaryDirectory() as root:
            _,state=granular_fixture(root)
            view=declared_view(state,needs_recheck=True,declared_state='confirmed')
            body=[p for p in view['pages'] if p['section']=='body']
            self.assertIn('재검토 필요 (이전 선언: 확인)',visible_blocks(body))
            summary=next(block for block in body[0]['blocks'] if block.get('variant')=='summary')
            self.assertEqual(summary['headers'],['구분','핵심 확인내용','판단','근거'])
            self.assertEqual(next(row[2] for row in summary['rows'] if row[0]=='주요 행위'),'재검토 필요 (이전 선언: 확인)')

    def test_same_evidence_labels_link_to_exact_source_appendix_in_both_formats(self):
        with tempfile.TemporaryDirectory() as root:
            _,state=granular_fixture(root); view=declared_view(state)
            mapping=labels(state)
            rendered=render_html(view)
            observed=state['observations'][0]
            target=bookmark_id(mapping[observed['id']])
            self.assertIn('id="'+target+'"',rendered)
            self.assertIn('href="#'+target+'"',rendered)
            output=Path(root)/'report.docx'; render_docx(view,output)
            with ZipFile(output) as archive:
                xml=archive.read('word/document.xml').decode()
                self.assertIn('w:name="'+target+'"',xml)
                self.assertIn('w:anchor="'+target+'"',xml)
                self.assertIn(observed['literal'],unescape(re.sub('<[^>]+>','',xml)))
            # Both formats consume each exact projected paragraph, not separate
            # prose generation paths or page-dependent answer summaries.
            for page in view['pages']:
                for block in page['blocks']:
                    if block['kind']=='paragraph':
                        self.assertIn(block['text'],unescape(re.sub('<[^>]+>','',rendered)))
                        # Hyperlink text is omitted by python-docx Paragraph.text.
                        self.assertIn(block['text'].replace('\n',''),unescape(re.sub('<[^>]+>','',xml)).replace('\n',''))

    def test_actual_bundle_records_style_and_export_identity_without_claiming_layout_review(self):
        with tempfile.TemporaryDirectory() as root:
            case,_,_=fixture(root);state=worked(case)
            result=bundle(case,state)
            manifest=json.loads(Path(result['manifest']).read_text())
            self.assertEqual(manifest['style_profile'],render_style_identity())
            self.assertEqual(manifest['layout_review']['status'],'not_rendered')
            self.assertIsNone(manifest['layout_review']['physical_page_count'])
            self.assertEqual(manifest['evidence_export']['status'],'complete')
            self.assertEqual({Path(p).name for p in result['files']},{'integrated.html'})
            archive=artifact(case,result['bundle_id'],'selected-evidence.zip')
            self.assertEqual(archive.name,'selected-evidence.zip')
            self.assertEqual(hashlib.sha256(archive.read_bytes()).hexdigest(),manifest['evidence_export']['sha256'])
            masked=bundle(case,state,redact=['maintenance'])
            self.assertEqual(masked['evidence_export']['status'],'not_exported_redacted')
            with self.assertRaises(ValueError):artifact(case,masked['bundle_id'],'selected-evidence.zip')

    def test_actual_common_style_profile_is_loaded_and_applied(self):
        with tempfile.TemporaryDirectory() as root:
            _,state=granular_fixture(root); view=declared_view(state)
            output=Path(root)/'styled.docx';render_docx(view,output)
            identity=render_style_identity()
            profile=Path(__file__).parent/'forsic_plugin/report_driven/report_style.json'
            self.assertEqual(identity['sha256'],hashlib.sha256(profile.read_bytes()).hexdigest())
            self.assertEqual(identity['source_kind'],'retained_docx_template')
            template=profile.parent/identity['template_file']
            self.assertEqual(identity['template_sha256'],hashlib.sha256(template.read_bytes()).hexdigest())
            self.assertEqual(json.loads(profile.read_text())['profile_id'],'forsic-integrated-template')
            document=Document(output); section=document.sections[0]
            for name,expected in [('top_margin',18),('bottom_margin',17),('left_margin',19),('right_margin',19),('header_distance',9),('footer_distance',10),('page_width',210),('page_height',297)]:
                self.assertAlmostEqual(getattr(section,name).mm,expected,delta=.03)
            self.assertEqual(document.styles['Normal'].font.size.pt,9.5)
            self.assertEqual(document.styles['Title'].font.size.pt,18.5)
            title_color=document.styles['Title'].font.color.rgb
            self.assertIsNotNone(title_color)
            self.assertEqual(len(set(title_color)),1)
            self.assertTrue(document.tables)
            for table in document.tables:
                if len(table.rows)==1:continue  # Template callout boxes have no header row.
                self.assertIsNotNone(table.rows[0]._tr.find(qn('w:trPr')).find(qn('w:tblHeader')))
                header=table.cell(0,0).paragraphs[0].runs[0]
                self.assertEqual(header.font.size.pt,8)
                self.assertEqual(str(header.font.color.rgb),'FFFFFF')
            with ZipFile(output) as archive:
                foot=archive.read('word/footer1.xml').decode()
                self.assertIn('PAGE',foot);self.assertNotIn('NUMPAGES',foot)


if __name__=='__main__':unittest.main()
